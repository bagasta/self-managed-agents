"""Persistent, owner-scoped chat with one bot or a group of bots.

Routing uses agent IDs selected by the composer. An unmentioned group request
goes to the manager, who can visibly delegate to specialists.
"""
from __future__ import annotations

import json
import re
import asyncio
import base64
import hashlib
import hmac
import shutil
import subprocess
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree

import structlog
from langchain_core.tools import tool
from fastapi import APIRouter, BackgroundTasks, Depends, File, Header, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from sqlalchemy import case, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.exc import IntegrityError

from app.api.workforce import WorkforcePrincipal, get_workforce_principal
from app.api.sessions import _ensure_arthur_ui_enterprise_test_owner, _ARTHUR_UI_ENTERPRISE_TEST_ID
from app.config import get_settings
from app.core.domain.agent_ownership import owner_filter
from app.core.domain.agent_quota_service import check_agent_quota, record_agent_token_usage
from app.core.domain.bot_profile import fill_missing_arthur_profile
from app.core.domain.product_evidence import looks_truncated, product_source_urls, research_result_has_sources
from app.core.domain.tenant_identity import resolve_unique_user
from app.core.domain.memory_service import get_active_context_version, get_versioned_memory, upsert_memory
from app.core.domain.skill_service import create_or_update_skill
from app.core.engine.agent_runner import run_agent
from app.core.engine.session_lock import session_run_lock
from app.database import AsyncSessionLocal, get_db
from app.models.agent import Agent
from app.models.session import Session
from app.models.team_chat import TeamChatAttachment, TeamChatMember, TeamChatMessage, TeamChatRoom
from app.models.subscription import User
from app.models.run import Run
from app.models.skill import Skill
from app.models.workforce_task import WorkforceTask, WorkforceTaskEvent, WorkforceTaskStep

router = APIRouter(prefix="/v1/team-chat", tags=["team-chat"])
log = structlog.get_logger(__name__)
MAX_GROUP_AUTOMATED_TURNS = 12
MAX_GROUP_TURNS_PER_AGENT = 2
TEAM_CHAT_UPLOAD_ROOT = Path("data/team-chat-uploads")
TEAM_CHAT_MAX_ATTACHMENT_BYTES = 8 * 1024 * 1024
TEAM_CHAT_MAX_IMAGE_BYTES = 5 * 1024 * 1024
TEAM_CHAT_ALLOWED_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".webp", ".gif", ".pdf", ".txt", ".md", ".csv", ".docx", ".pptx", ".xlsx",
}
TEAM_CHAT_IMAGE_TYPES = {"image/png", "image/jpeg", "image/webp", "image/gif"}
TEAM_CHAT_ATTACHMENT_PREVIEW_CHARS = 6_000
COMPUTER_FILE_TOKEN_TTL_SECONDS = 15 * 60


async def get_team_chat_principal(
    request: Request,
    x_api_key: str | None = Header(None, alias="X-API-Key"),
    x_user_key: str | None = Header(None, alias="X-User-Key"),
    db: AsyncSession = Depends(get_db),
) -> WorkforcePrincipal:
    if x_api_key or x_user_key:
        principal = await get_workforce_principal(x_api_key=x_api_key, x_user_key=x_user_key, db=db)
        if principal.is_platform_admin and request.query_params.get("owner_view") == "true":
            workspace = (request.query_params.get("workspace_id") or "").strip()
            if not workspace:
                raise HTTPException(422, "Pilih akun Owner untuk Team Chat")
            if workspace == _ARTHUR_UI_ENTERPRISE_TEST_ID:
                if get_settings().environment.lower() not in {"development", "dev", "test"}:
                    raise HTTPException(403, "Akun QA hanya tersedia di development")
                await _ensure_arthur_ui_enterprise_test_owner(db)
            owner = await resolve_unique_user(db, workspace)
            if owner is None:
                raise HTTPException(404, "Akun Owner tidak ditemukan atau tidak unik")
            return WorkforcePrincipal(
                owner_user_id=owner.id,
                owner_external_id=owner.external_id,
                is_platform_admin=True,
            )
        return principal
    # This dashboard is a local development harness. Never grant this implicit
    # test identity through a public host, reverse proxy, or production mode.
    host = request.url.hostname
    client_host = request.client.host if request.client else None
    origin = request.headers.get("origin")
    referer = request.headers.get("referer")
    same_origin = f"{request.url.scheme}://{request.headers.get('host', '')}"
    if (get_settings().environment.lower() not in {"development", "dev", "test"}
        or host not in {"127.0.0.1", "localhost", "::1"}
        or client_host not in {"127.0.0.1", "::1"}
        or (origin and origin != same_origin)
        or (referer and not referer.startswith(same_origin + "/ui/"))):
        raise HTTPException(401, "A platform or owner-bound API key is required")
    await _ensure_arthur_ui_enterprise_test_owner(db)
    owner = (await db.execute(select(User).where(User.external_id == _ARTHUR_UI_ENTERPRISE_TEST_ID))).scalar_one()
    return WorkforcePrincipal(owner_user_id=owner.id, owner_external_id=_ARTHUR_UI_ENTERPRISE_TEST_ID)


class RoomCreate(BaseModel):
    kind: str = Field(pattern="^(direct|group)$")
    title: str = Field(min_length=1, max_length=255)
    manager_agent_id: uuid.UUID
    member_agent_ids: list[uuid.UUID] = Field(default_factory=list, max_length=20)


class ChatSend(BaseModel):
    content: str = Field(default="", max_length=16000)
    target_agent_ids: list[uuid.UUID] = Field(default_factory=list, max_length=20)
    attachment_ids: list[uuid.UUID] = Field(default_factory=list, max_length=5)
    everyone: bool = False
    client_message_id: uuid.UUID | None = None


class BotConfigUpdate(BaseModel):
    name: str = Field(min_length=1, max_length=255)
    description: str = Field(default="", max_length=20_000)
    instructions: str = Field(default="", max_length=200_000)
    identity: str = Field(default="", max_length=30_000)
    soul: str = Field(default="", max_length=30_000)
    model: str = Field(min_length=1, max_length=255)
    temperature: float = Field(ge=0, le=2)
    skills_enabled: bool = True
    # Optional keeps older clients from silently revoking an existing computer
    # assignment when they save an unrelated profile change.
    computer_enabled: bool | None = None


class BotSkillWrite(BaseModel):
    name: str = Field(min_length=1, max_length=255)
    description: str = Field(min_length=1, max_length=20_000)
    content_md: str = Field(min_length=1, max_length=100_000)


def _workspace(principal: WorkforcePrincipal, requested: str | None) -> str:
    if principal.owner_user_id:
        return str(principal.owner_user_id)
    if not requested:
        raise HTTPException(422, "workspace_id is required for platform admin access")
    return requested


async def _roster(db: AsyncSession, principal: WorkforcePrincipal, workspace: str) -> list[Agent]:
    scope = Agent.owner_user_id == principal.owner_user_id if principal.owner_user_id else owner_filter(workspace)
    agents = list((await db.execute(select(Agent).where(Agent.is_deleted.is_(False), scope).order_by(Agent.created_at))).scalars().all())
    arthur = (await db.execute(select(Agent).where(Agent.is_deleted.is_(False), Agent.tools_config.contains({"system_plugin": "arthur_v2"})).order_by(Agent.created_at.desc()).limit(1))).scalar_one_or_none()
    if arthur and all(agent.id != arthur.id for agent in agents):
        agents.insert(0, arthur)
    return agents


async def _configurable_agent(
    db: AsyncSession, principal: WorkforcePrincipal, workspace: str, agent_id: uuid.UUID,
) -> Agent:
    agent = (await db.execute(select(Agent).where(Agent.id == agent_id, Agent.is_deleted.is_(False)))).scalar_one_or_none()
    if agent is None:
        raise HTTPException(404, "Bot tidak ditemukan")
    owned = agent.owner_user_id == principal.owner_user_id if principal.owner_user_id else agent in await _roster(db, principal, workspace)
    local_arthur = (
        principal.owner_external_id == _ARTHUR_UI_ENTERPRISE_TEST_ID
        and (agent.tools_config or {}).get("system_plugin") == "arthur_v2"
    )
    if not (owned or local_arthur):
        raise HTTPException(404, "Bot tidak ditemukan di workspace ini")
    return agent


async def _bot_config_data(db: AsyncSession, agent: Agent) -> dict:
    identity = await get_versioned_memory(agent.id, "identity", db)
    soul = await get_versioned_memory(agent.id, "soul", db)
    skills = (await db.execute(select(Skill).where(
        Skill.agent_id == agent.id, Skill.enabled.is_(True),
    ).order_by(Skill.name, Skill.updated_at.desc()))).scalars().all()
    return {
        "id": agent.id, "name": agent.name, "description": agent.description or "",
        "instructions": agent.instructions or "", "identity": identity.value_data if identity else "",
        "soul": soul.value_data if soul else "", "model": agent.model,
        "temperature": agent.temperature,
        "skills_enabled": (agent.tools_config or {}).get("skills", True),
        "computer_enabled": bool((agent.tools_config or {}).get("computer", False)),
        "skills": [
            {"name": skill.name, "description": skill.description,
             "content_md": skill.content_md, "editable": not skill.immutable}
            for skill in skills
        ],
    }


def _room_data(room: TeamChatRoom, members: list[TeamChatMember]) -> dict:
    return {"id": room.id, "kind": room.kind, "title": room.title,
            "manager_agent_id": room.manager_agent_id,
            "member_agent_ids": [member.agent_id for member in members],
            "updated_at": room.updated_at}


def _attachment_data(attachment: TeamChatAttachment) -> dict:
    return {
        "id": attachment.id, "filename": attachment.filename,
        "content_type": attachment.content_type, "size_bytes": attachment.size_bytes,
        "is_image": attachment.content_type in TEAM_CHAT_IMAGE_TYPES,
        "download_url": f"/v1/team-chat/rooms/{attachment.room_id}/attachments/{attachment.id}/download",
    }


def _message_data(message: TeamChatMessage, attachments: list[TeamChatAttachment] | None = None) -> dict:
    return {"id": message.id, "sender_type": message.sender_type,
            "sender_agent_id": message.sender_agent_id, "content": message.content,
            "status": message.status, "created_at": message.created_at,
            "attachments": [_attachment_data(item) for item in attachments or []]}


def _safe_attachment_filename(raw: str | None) -> str:
    name = Path(raw or "attachment").name.strip()
    return (name or "attachment")[:255]


def _attachment_path(attachment: TeamChatAttachment) -> Path:
    return TEAM_CHAT_UPLOAD_ROOT / attachment.storage_key


def _computer_file_token(*, attachment: TeamChatAttachment, room: TeamChatRoom, principal: WorkforcePrincipal) -> str:
    """Create a short-lived, owner-scoped URL token for the isolated browser VM."""
    claims = {
        "attachment_id": str(attachment.id),
        "room_id": str(room.id),
        "owner_id": str(principal.owner_user_id or ""),
        "exp": int(datetime.now(timezone.utc).timestamp()) + COMPUTER_FILE_TOKEN_TTL_SECONDS,
    }
    encoded = base64.urlsafe_b64encode(json.dumps(claims, separators=(",", ":")).encode()).decode().rstrip("=")
    key = get_settings().api_key.encode()
    signature = hmac.new(key, encoded.encode(), hashlib.sha256).hexdigest()
    return f"{encoded}.{signature}"


def _verified_computer_file_claims(token: str, attachment_id: uuid.UUID) -> dict | None:
    try:
        encoded, signature = token.rsplit(".", 1)
        expected = hmac.new(get_settings().api_key.encode(), encoded.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected):
            return None
        payload = json.loads(base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4)))
        if (not isinstance(payload, dict)
                or payload.get("attachment_id") != str(attachment_id)
                or not isinstance(payload.get("exp"), int)
                or payload["exp"] < int(datetime.now(timezone.utc).timestamp())):
            return None
        return payload
    except (ValueError, TypeError, json.JSONDecodeError):
        return None


def _computer_attachment_url(attachment: TeamChatAttachment, room: TeamChatRoom, principal: WorkforcePrincipal) -> str | None:
    settings = get_settings()
    base_url = (getattr(settings, "computer_runtime_file_base_url", "") or settings.app_public_url or "").rstrip("/")
    if not base_url:
        return None
    token = _computer_file_token(attachment=attachment, room=room, principal=principal)
    return f"{base_url}/v1/team-chat/computer-files/{attachment.id}?token={token}"


async def _workforce_dispatch_receipt(
    db: AsyncSession, *, steps: list[dict], workspace: str, owner_user_id: uuid.UUID | None,
) -> tuple[str, list[WorkforceTaskStep]]:
    """Show only assignments backed by a successful workforce tool call and owned task."""
    task_ids: list[uuid.UUID] = []
    for step in steps:
        if step.get("tool") != "orchestrate_owner_workforce_task":
            continue
        try:
            raw = step.get("result")
            result = json.loads(raw) if isinstance(raw, str) else raw
            if isinstance(result, dict) and result.get("ok") is True:
                task_id = uuid.UUID(str(result["task_id"]))
                if task_id not in task_ids:
                    task_ids.append(task_id)
        except (KeyError, TypeError, ValueError):
            continue
    if not task_ids:
        return "", []

    summaries: list[str] = []
    dispatched_steps: list[WorkforceTaskStep] = []
    for task_id in task_ids:
        scope = [WorkforceTask.id == task_id, WorkforceTask.workspace_id == workspace]
        if owner_user_id:
            scope.append(WorkforceTask.owner_user_id == owner_user_id)
        task = (await db.execute(select(WorkforceTask).where(*scope))).scalar_one_or_none()
        if task is None:
            continue
        assignments = (await db.execute(
            select(WorkforceTaskStep, Agent.name)
            .join(Agent, Agent.id == WorkforceTaskStep.assigned_agent_id)
            .where(WorkforceTaskStep.task_id == task.id)
            .order_by(WorkforceTaskStep.sequence)
        )).all()
        names = ", ".join(f"@{name}" for _, name in assignments)
        if names:
            if task.status == "blocked":
                summaries.append(f"Aku sudah membagikan {task.title} ke {names}. Ada bagian yang memerlukan tindak lanjut sebelum hasil lengkap tersedia.")
            elif task.status == "completed":
                summaries.append(f"{names} sudah mengerjakan {task.title}. Hasilnya ada di balasan mereka di bawah.")
            else:
                summaries.append(f"Aku sudah minta {names} mengerjakan {task.title}. Hasilnya akan muncul di grup ini.")
        for assignment, _ in assignments:
            dispatched_steps.append(assignment)
    return "\n\n".join(summaries), dispatched_steps


async def _workforce_status_receipt(
    db: AsyncSession, *, steps: list[dict], workspace: str, owner_user_id: uuid.UUID | None,
) -> str:
    """Present a roster read from persisted tasks, never from model-authored metadata."""
    if not any(step.get("tool") == "list_owner_workforce_tasks" for step in steps):
        return ""
    scope = [WorkforceTask.workspace_id == workspace]
    if owner_user_id:
        scope.append(WorkforceTask.owner_user_id == owner_user_id)
    tasks = (await db.execute(
        select(WorkforceTask).where(*scope).order_by(WorkforceTask.created_at.desc()).limit(5)
    )).scalars().all()
    lines: list[str] = []
    for task in tasks:
        names = (await db.execute(
            select(Agent.name).join(WorkforceTaskStep, WorkforceTaskStep.assigned_agent_id == Agent.id)
            .where(WorkforceTaskStep.task_id == task.id).order_by(WorkforceTaskStep.sequence)
        )).scalars().all()
        team = ", ".join(f"@{name}" for name in dict.fromkeys(names)) or "Tim"
        if task.status == "completed":
            lines.append(f"{team} sudah menyelesaikan {task.title}.")
        elif task.status in {"blocked", "failed"}:
            lines.append(f"{team} belum bisa menyelesaikan {task.title}. Kendalanya perlu ditangani sebelum pekerjaan dilanjutkan.")
        elif task.status == "cancelled":
            lines.append(f"Pekerjaan {task.title} sudah dihentikan.")
        else:
            lines.append(f"{team} sedang mengerjakan {task.title}.")
    return "Ini kabar terbaru dari tim:\n\n" + "\n".join(f"• {line}" for line in lines) if lines else "Belum ada pekerjaan tim yang tercatat."


def _requested_workforce_status(steps: list[dict]) -> bool:
    """A manager asked the durable workforce control plane for a team update."""
    return any(step.get("tool") == "list_owner_workforce_tasks" for step in steps)


def _group_message_tool(member_names: dict[str, uuid.UUID], sender_id: uuid.UUID):
    @tool
    async def message_group_members(agent_names: list[str], message: str) -> str:
        """Address other bots in this group. Give their names and the message they should receive. Their replies will appear in the group chat."""
        names = [name.strip().casefold() for name in agent_names]
        targets = list(dict.fromkeys(member_names[name] for name in names if name in member_names))
        targets = [agent_id for agent_id in targets if agent_id != sender_id]
        if not targets or not message.strip():
            return json.dumps({"ok": False, "error": "Choose group member names and a nonempty message."})
        return json.dumps({"ok": True, "agent_ids": [str(agent_id) for agent_id in targets], "message": message.strip()[:4000]})

    return message_group_members


def _group_message_requests(steps: list[dict], members: set[uuid.UUID], sender_id: uuid.UUID) -> list[tuple[list[uuid.UUID], str]]:
    requests = []
    for step in steps:
        if step.get("tool") != "message_group_members":
            continue
        try:
            raw = step.get("result")
            result = json.loads(raw) if isinstance(raw, str) else raw
            if not isinstance(result, dict) or result.get("ok") is not True:
                continue
            ids = list(dict.fromkeys(uuid.UUID(value) for value in result.get("agent_ids", [])))
            message = str(result.get("message") or "").strip()
            if ids and len(ids) <= 20 and all(agent_id in members and agent_id != sender_id for agent_id in ids) and message:
                requests.append((ids, message[:4000]))
        except (TypeError, ValueError, AttributeError):
            continue
    return requests


def _reply_handoff_targets(reply: str, names: dict[uuid.UUID, str], sender_id: uuid.UUID) -> list[uuid.UUID]:
    """Every explicit mention of another room member starts a bot-to-bot turn."""
    matches: list[tuple[int, uuid.UUID]] = []
    for agent_id, name in names.items():
        if agent_id == sender_id or not name:
            continue
        mention = re.search(r"(?<![\w@])@" + re.escape(name) + r"(?!\w)", reply or "", re.IGNORECASE)
        if mention:
            matches.append((mention.start(), agent_id))
    return [agent_id for _, agent_id in sorted(matches, key=lambda match: match[0])]


def _implicit_group_target(content: str, names: dict[uuid.UUID, str], manager_id: uuid.UUID) -> uuid.UUID:
    """A group request without an explicit recipient always reaches its manager."""
    return manager_id


def _correct_group_reply_recipient(reply: str, manager_name: str, actual_sender: str | None) -> str:
    """Remove a stale manager salutation only on replies to the owner."""
    if not manager_name or actual_sender is not None:
        return reply
    prefix = re.compile(r"^\s*@" + re.escape(manager_name) + r"(?=$|[\s,:-])\s*[:,\-]?\s*", re.IGNORECASE)
    return prefix.sub("", reply or "", count=1)


_OWNER_ADDRESS = re.compile(
    r"^\s*(?:(?:untuk|kepada)\s+)?(?:user|owner|pemilik|bos|boss|pak|bu)\b\s*[:,\-]",
    re.IGNORECASE,
)


def _address_bot_reply(reply: str, sender_id: uuid.UUID, member_names: dict[uuid.UUID, str]) -> str:
    """Make a bot-directed answer address its bot by mention, not bare name."""
    text = reply or ""
    sender_name = member_names.get(sender_id, "")
    if not text.strip() or not sender_name or _OWNER_ADDRESS.match(text):
        return text
    names = sorted((name for name in member_names.values() if name), key=len, reverse=True)
    for name in names:
        if re.match(r"^\s*@" + re.escape(name) + r"(?!\w)", text, re.IGNORECASE):
            return text
        for marker in ("**", "__", ""):
            bare = re.match(
                r"^(\s*)" + re.escape(marker + name + marker) + r"(?=\s*[:,—–-])",
                text, re.IGNORECASE,
            )
            if bare:
                return bare.group(1) + "@" + name + text[bare.end():]
    if _reply_handoff_targets(text, member_names, sender_id):
        return text
    return f"@{sender_name} {text.lstrip()}"


_DEFERRED_WORK_PROMISE = re.compile(
    r"\b(?:begitu\s+(?:sudah\s+)?jadi|"
    r"(?:nanti|setelah\s+(?:selesai|jadi))\b.{0,70}\b(?:kirim|post|posting|bagikan|share)|"
    r"(?:target|selesai|rampung)\b.{0,110}\b(?:dalam\s+)?\d+\s*(?:[-–]\s*\d+\s*)?(?:jam|hari))",
    re.IGNORECASE | re.DOTALL,
)


def _unbacked_deferred_work(reply: str, steps: list[dict]) -> bool:
    """A chat promise is not a background job or the requested deliverable."""
    for step in steps:
        if step.get("tool") != "orchestrate_owner_workforce_task":
            continue
        try:
            raw = step.get("result")
            result = json.loads(raw) if isinstance(raw, str) else raw
            if isinstance(result, dict) and result.get("ok") is True:
                uuid.UUID(str(result["task_id"]))
                return False
        except (KeyError, TypeError, ValueError):
            continue
    text = reply or ""
    if not _DEFERRED_WORK_PROMISE.search(text):
        return False
    # A substantial artifact may also describe a later review or publishing step.
    return len(text.strip()) < 1000


_UNVERIFIED_GROUP_DATA_CLAIM = re.compile(
    r"\b(?:data(?:nya)?|hasil(?:\s+riset)?|tabel)\s+(?:(?:sudah|udah|telah)\s+)?(?:masuk|diterima|lengkap|siap)\b",
    re.IGNORECASE,
)
_GROUP_RESEARCH_TOOL_NAMES = frozenset({"tavily_search", "tavily_extract", "http_get"})


def _claims_product_prices(reply: str) -> bool:
    """A blocker asking a teammate for help is not a completed price report."""
    if re.search(r"\bRp\s*[\d.,]+", reply or "", re.IGNORECASE):
        return True
    return bool(
        re.search(r"\b(?:tabel|daftar)\s+harga\b", reply or "", re.IGNORECASE)
        and not re.search(r"\b(?:belum|gagal|pending|tidak bisa)\b", reply or "", re.IGNORECASE)
    )


def _has_group_research_evidence(steps: list[dict], message_context: str) -> bool:
    """Return true only for visible group data or a successful research-tool result."""
    product_links = product_source_urls(message_context)
    counts = [int(value) for value in re.findall(r"\b(\d{1,2})\s+produk\b", message_context, re.IGNORECASE)]
    required = max(counts, default=2)
    if len(product_links) >= required:
        return True
    tool_evidence = set(product_links)
    for step in steps:
        if step.get("tool") not in _GROUP_RESEARCH_TOOL_NAMES:
            continue
        result = str(step.get("result") or "").strip()
        if result and not result.casefold().startswith("[error]"):
            tool_evidence.update(product_source_urls(result))
    return len(tool_evidence) >= required


def guard_group_reply_evidence(reply: str, steps: list[dict], message_context: str) -> str:
    """Reject a claim that research data arrived when no chat/tool evidence exists."""
    quantity_claim = re.search(r"\b(?:data\s+)?(\d{1,2})\s+produk\b.{0,35}\b(?:siap|lengkap|masuk|selesai)\b", reply or "", re.IGNORECASE)
    if not _UNVERIFIED_GROUP_DATA_CLAIM.search(reply or "") and not quantity_claim:
        return reply
    evidence_context = f"{message_context}\n{quantity_claim.group(1)} produk" if quantity_claim else message_context
    if _has_group_research_evidence(steps, evidence_context):
        return reply
    return (
        "Saya belum menerima data riset yang terverifikasi di percakapan ini, jadi belum bisa "
        "menyatakan data atau tabel sudah lengkap. Saya menunggu tabel atau tautan sumber yang nyata."
    )


async def _agent_task_snapshot(
    db: AsyncSession, *, agent_id: uuid.UUID, workspace: str,
    owner_user_id: uuid.UUID | None,
) -> str:
    scope = [
        WorkforceTask.workspace_id == workspace,
        WorkforceTaskStep.assigned_agent_id == agent_id,
        WorkforceTaskStep.parent_step_id.is_(None),
    ]
    if owner_user_id:
        scope.append(WorkforceTask.owner_user_id == owner_user_id)
    rows = (await db.execute(
        select(WorkforceTaskStep, WorkforceTask)
        .join(WorkforceTask, WorkforceTask.id == WorkforceTaskStep.task_id)
        .where(*scope)
        .order_by(WorkforceTask.created_at.desc(), WorkforceTaskStep.sequence)
        .limit(3)
    )).all()
    if not rows:
        return "No assignments are recorded for you in this workspace. Do not invent a current task or completion."
    return "\n".join(
        f"- {step.title} ({task.title}): {step.status}. "
        f"Recorded result: {(step.output_summary or 'No result recorded yet')[:320]}"
        for step, task in rows
    )


async def _sync_workforce_replies(db: AsyncSession, room: TeamChatRoom) -> None:
    # A bot removed from the roster cannot finish an old placeholder.
    retired = (await db.execute(
        select(TeamChatMessage)
        .join(Agent, Agent.id == TeamChatMessage.sender_agent_id)
        .where(
            TeamChatMessage.room_id == room.id,
            TeamChatMessage.sender_type == "agent",
            TeamChatMessage.status.in_(["working", "failed"]),
            Agent.is_deleted.is_(True),
        )
    )).scalars().all()
    for message in retired:
        if not message.content:
            message.content = "Bot ini sudah tidak aktif lagi."
        message.status = "failed"
    pending = (await db.execute(select(TeamChatMessage).where(
        TeamChatMessage.room_id == room.id,
        TeamChatMessage.workforce_step_id.is_not(None),
        TeamChatMessage.status == "working",
    ))).scalars().all()
    if not pending:
        if retired:
            await db.commit()
        return
    step_ids = [message.workforce_step_id for message in pending]
    scope = [WorkforceTaskStep.id.in_(step_ids), WorkforceTask.workspace_id == room.workspace_id]
    if room.owner_user_id:
        scope.append(WorkforceTask.owner_user_id == room.owner_user_id)
    steps = (await db.execute(
        select(WorkforceTaskStep).join(WorkforceTask, WorkforceTask.id == WorkforceTaskStep.task_id).where(*scope)
    )).scalars().all()
    by_id = {step.id: step for step in steps}
    run_ids = [step.run_id for step in steps if step.run_id]
    runs = (await db.execute(select(Run).where(Run.id.in_(run_ids)))).scalars().all() if run_ids else []
    by_run_id = {run.id: run for run in runs}
    working_ids = [message.workforce_step_id for message in pending if message.status == "working"]
    progress_events = (await db.execute(select(WorkforceTaskEvent).where(
        WorkforceTaskEvent.step_id.in_(working_ids),
        WorkforceTaskEvent.event_type == "team_message",
    ).order_by(WorkforceTaskEvent.created_at.desc(), WorkforceTaskEvent.id.desc()).limit(200))).scalars().all() if working_ids else []
    changed = False
    pending_by_step = {message.workforce_step_id: message for message in pending}
    for event in reversed(progress_events):
        placeholder = pending_by_step.get(event.step_id)
        content = str((event.metadata_ or {}).get("message") or "").strip()[:4000]
        if not placeholder or not content:
            continue
        progress_id = uuid.uuid5(uuid.NAMESPACE_URL, f"team-chat-workforce-event:{event.id}")
        if await db.get(TeamChatMessage, progress_id) is None:
            db.add(TeamChatMessage(
                id=progress_id, room_id=room.id, sender_type="agent",
                sender_agent_id=placeholder.sender_agent_id, content=content,
                status="sent", created_at=event.created_at,
            ))
            changed = True
    for message in pending:
        step = by_id.get(message.workforce_step_id)
        if step is None:
            continue
        if step.status in {"pending", "queued", "in_progress"}:
            continue
        if step.status == "completed":
            result = (step.output_summary or "Aku sudah menyelesaikan tugas ini.").strip()
            content = result
            status = "sent"
        elif step.status == "cancelled":
            content, status = "Tugas ini dihentikan.", "failed"
        else:
            run = by_run_id.get(step.run_id)
            failure_code = (run.runtime_metadata or {}).get("failure_code") if run else None
            if failure_code in {"sandbox_runtime_unavailable", "sandbox_capability_disabled"}:
                content = "Bagian ini memerlukan lingkungan kerja yang belum tersedia."
            elif failure_code == "unverified_product_sources":
                content = "Riset belum dapat diselesaikan karena sumber produk belum terverifikasi."
            else:
                content = "Belum ada hasil yang dapat ditampilkan untuk bagian ini."
            status = "failed"
        if message.content:
            # An already displayed progress message is immutable. Close its
            # spinner and append the result instead of rewriting history.
            message.status = "sent"
            result_id = uuid.uuid5(uuid.NAMESPACE_URL, f"team-chat-workforce-result:{step.id}")
            if await db.get(TeamChatMessage, result_id) is None:
                db.add(TeamChatMessage(
                    id=result_id, room_id=room.id, sender_type="agent",
                    sender_agent_id=message.sender_agent_id, content=content,
                    status=status, created_at=datetime.now(timezone.utc),
                ))
            changed = True
        else:
            message.content, message.status = content, status
            message.created_at = datetime.now(timezone.utc)
            changed = True
    if changed:
        await db.commit()
    elif retired:
        await db.commit()


async def _room_members(db: AsyncSession, room_id: uuid.UUID) -> list[TeamChatMember]:
    # Keep soft-deleted bots out of the live room roster. Old rooms can outlive
    # their bots; treating those rows as active creates permanent typing bubbles
    # and dispatch attempts that fail before a Run can even be recorded.
    return (await db.execute(
        select(TeamChatMember)
        .join(Agent, Agent.id == TeamChatMember.agent_id)
        .where(TeamChatMember.room_id == room_id, Agent.is_deleted.is_(False))
    )).scalars().all()


async def _ensure_member_session(
    db: AsyncSession, room: TeamChatRoom, member: TeamChatMember,
    principal: WorkforcePrincipal,
) -> None:
    agent = await db.get(Agent, member.agent_id)
    session = await db.get(Session, member.session_id) if member.session_id else None
    if not agent or agent.is_deleted:
        return
    if session and session.agent_id == agent.id:
        return
    metadata = {"team_chat_room_id": str(room.id), "kind": room.kind}
    if principal.owner_external_id == _ARTHUR_UI_ENTERPRISE_TEST_ID:
        metadata.update({"source": "arthur-ui", "memory_mode": "isolated", "test_plan": "enterprise"})
    session = Session(
        agent_id=agent.id,
        external_user_id=principal.owner_external_id or room.workspace_id,
        metadata_=metadata,
        channel_type=None,
    )
    db.add(session)
    await db.flush()
    member.session_id = session.id


async def _require_room(db: AsyncSession, principal: WorkforcePrincipal, workspace: str, room_id: uuid.UUID) -> TeamChatRoom:
    room = (await db.execute(select(TeamChatRoom).where(
        TeamChatRoom.id == room_id, TeamChatRoom.workspace_id == workspace,
        TeamChatRoom.kind.in_(["direct", "group"]),
    ))).scalar_one_or_none()
    if not room or (principal.owner_user_id and room.owner_user_id != principal.owner_user_id):
        raise HTTPException(404, "Chat room not found")
    return room


@router.get("/roster")
async def roster(workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    workspace = _workspace(principal, workspace_id)
    agents = await _roster(db, principal, workspace)
    active = set((await db.execute(
        select(Session.agent_id).join(Run, Run.session_id == Session.id)
        .join(TeamChatMember, TeamChatMember.session_id == Session.id)
        .join(TeamChatRoom, TeamChatRoom.id == TeamChatMember.room_id)
        .where(TeamChatRoom.workspace_id == workspace, Run.status == "running")
        .distinct()
    )).scalars())
    active.update((await db.execute(
        select(WorkforceTaskStep.assigned_agent_id)
        .join(WorkforceTask, WorkforceTask.id == WorkforceTaskStep.task_id)
        .where(WorkforceTask.workspace_id == workspace, WorkforceTaskStep.status == "in_progress", WorkforceTask.status == "in_progress")
        .distinct()
    )).scalars())
    return {"workspace_id": workspace, "items": [{"id": a.id, "name": a.name, "description": a.description,
            "is_manager": bool((a.tools_config or {}).get("builder") or (a.tools_config or {}).get("wa_agent_manager") or "builder" in (a.capabilities or [])),
            "is_arthur": (a.tools_config or {}).get("system_plugin") == "arthur_v2",
            "is_working": a.id in active,
            "computer_enabled": bool((a.tools_config or {}).get("computer", False))} for a in agents]}


@router.get("/agents/{agent_id}/config")
async def get_bot_config(agent_id: uuid.UUID, workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    agent = await _configurable_agent(db, principal, _workspace(principal, workspace_id), agent_id)
    if await fill_missing_arthur_profile(agent, db):
        await db.commit()
    return await _bot_config_data(db, agent)


@router.patch("/agents/{agent_id}/config")
async def update_bot_config(agent_id: uuid.UUID, payload: BotConfigUpdate, workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    workspace = _workspace(principal, workspace_id)
    agent = await _configurable_agent(db, principal, workspace, agent_id)
    old_name = agent.name
    agent.name = payload.name.strip()
    if not agent.name:
        raise HTTPException(422, "Nama bot wajib diisi")
    agent.description = payload.description.strip() or None
    agent.instructions = payload.instructions.strip()
    agent.model = payload.model.strip()
    if not agent.model:
        raise HTTPException(422, "Model wajib diisi")
    agent.temperature = payload.temperature
    tools_config = {**(agent.tools_config or {}), "skills": payload.skills_enabled}
    if payload.computer_enabled is not None:
        tools_config["computer"] = payload.computer_enabled
    agent.tools_config = tools_config
    agent.version += 1
    context_version = await get_active_context_version(agent.id, db)
    for key, value in (("identity", payload.identity), ("soul", payload.soul)):
        memory_key = f"{key}:v{context_version}" if context_version else key
        await upsert_memory(agent.id, memory_key, value.strip(), db)
    if old_name != agent.name:
        direct_rooms = (await db.execute(select(TeamChatRoom).where(
            TeamChatRoom.workspace_id == workspace, TeamChatRoom.kind == "direct",
            TeamChatRoom.manager_agent_id == agent.id, TeamChatRoom.title == old_name,
        ))).scalars().all()
        for direct_room in direct_rooms:
            direct_room.title = agent.name
    await db.commit()
    return await _bot_config_data(db, agent)


@router.post("/agents/{agent_id}/skills")
async def save_bot_skill(agent_id: uuid.UUID, payload: BotSkillWrite, workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    agent = await _configurable_agent(db, principal, _workspace(principal, workspace_id), agent_id)
    name, description, content = payload.name.strip(), payload.description.strip(), payload.content_md.strip()
    if not name or not description or not content:
        raise HTTPException(422, "Nama, deskripsi, dan isi skill wajib diisi")
    system_skill = (await db.execute(select(Skill.id).where(
        Skill.agent_id == agent.id, Skill.name == name, Skill.immutable.is_(True),
    ).limit(1))).scalar_one_or_none()
    if system_skill:
        raise HTTPException(409, "Skill bawaan tidak dapat ditimpa dari panel ini")
    await create_or_update_skill(agent.id, name, description, content, db)
    await db.commit()
    return await _bot_config_data(db, agent)


@router.delete("/agents/{agent_id}/skills/{name}")
async def remove_bot_skill(agent_id: uuid.UUID, name: str, workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    agent = await _configurable_agent(db, principal, _workspace(principal, workspace_id), agent_id)
    skill = (await db.execute(select(Skill).where(
        Skill.agent_id == agent.id, Skill.name == name, Skill.version == "user",
        Skill.immutable.is_(False),
    ))).scalar_one_or_none()
    if skill is None:
        raise HTTPException(404, "Skill yang bisa diedit tidak ditemukan")
    await db.delete(skill)
    await db.commit()
    return await _bot_config_data(db, agent)


@router.get("/rooms")
async def list_rooms(workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    workspace = _workspace(principal, workspace_id)
    rooms = (await db.execute(select(TeamChatRoom).where(TeamChatRoom.workspace_id == workspace, TeamChatRoom.kind.in_(["direct", "group"]), *([TeamChatRoom.owner_user_id == principal.owner_user_id] if principal.owner_user_id else [])).order_by(TeamChatRoom.updated_at.desc()))).scalars().all()
    active_ids = {agent.id for agent in await _roster(db, principal, workspace)}
    items = []
    for room in rooms:
        if room.kind == "direct" and room.manager_agent_id not in active_ids:
            continue
        items.append(_room_data(room, await _room_members(db, room.id)))
    return {"items": items}


@router.post("/rooms", status_code=201)
async def create_room(payload: RoomCreate, workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    workspace = _workspace(principal, workspace_id)
    allowed = {a.id: a for a in await _roster(db, principal, workspace)}
    ids = list(dict.fromkeys([payload.manager_agent_id, *payload.member_agent_ids]))
    if not ids or any(agent_id not in allowed for agent_id in ids):
        raise HTTPException(422, "Every member must belong to this workspace")
    if payload.kind == "direct" and len(ids) != 1:
        raise HTTPException(422, "Direct chat has exactly one bot")
    if payload.kind == "group" and len(ids) < 2:
        raise HTTPException(422, "Group chat needs at least two bots")
    if payload.kind == "direct":
        existing = (await db.execute(select(TeamChatRoom).where(TeamChatRoom.workspace_id == workspace, TeamChatRoom.kind == "direct", TeamChatRoom.manager_agent_id == payload.manager_agent_id))).scalar_one_or_none()
        if existing:
            return _room_data(existing, await _room_members(db, existing.id))
    room = TeamChatRoom(workspace_id=workspace, owner_user_id=principal.owner_user_id, kind=payload.kind,
                        title=payload.title.strip(), manager_agent_id=payload.manager_agent_id)
    db.add(room)
    await db.flush()
    members = []
    for agent_id in ids:
        local_test = principal.owner_external_id == _ARTHUR_UI_ENTERPRISE_TEST_ID
        metadata = {"team_chat_room_id": str(room.id), "kind": payload.kind}
        if local_test:
            metadata.update({"source": "arthur-ui", "memory_mode": "isolated", "test_plan": "enterprise"})
        session = Session(agent_id=agent_id, external_user_id=principal.owner_external_id or workspace,
                          metadata_=metadata, channel_type=None)
        db.add(session)
        await db.flush()
        member = TeamChatMember(room_id=room.id, agent_id=agent_id, session_id=session.id)
        db.add(member)
        members.append(member)
    await db.commit()
    return _room_data(room, members)


@router.get("/rooms/{room_id}/messages")
async def list_messages(room_id: uuid.UUID, workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    room = await _require_room(db, principal, _workspace(principal, workspace_id), room_id)
    await _sync_workforce_replies(db, room)
    rows = (await db.execute(select(TeamChatMessage).where(TeamChatMessage.room_id == room.id).order_by(
        TeamChatMessage.created_at.desc(),
        case((TeamChatMessage.sender_type == "owner", 0), else_=1).desc(),
        TeamChatMessage.id.desc(),
    ).limit(200))).scalars().all()
    message_ids = [row.id for row in rows]
    attachment_rows = (await db.execute(select(TeamChatAttachment).where(
        TeamChatAttachment.message_id.in_(message_ids),
    ).order_by(TeamChatAttachment.created_at))).scalars().all() if message_ids else []
    attachments_by_message: dict[uuid.UUID, list[TeamChatAttachment]] = {}
    for attachment in attachment_rows:
        if attachment.message_id:
            attachments_by_message.setdefault(attachment.message_id, []).append(attachment)
    return {"items": [_message_data(row, attachments_by_message.get(row.id)) for row in reversed(rows)]}


@router.post("/rooms/{room_id}/attachments", status_code=201)
async def upload_attachment(
    room_id: uuid.UUID,
    file: UploadFile = File(...),
    workspace_id: str | None = Query(None, max_length=64),
    db: AsyncSession = Depends(get_db),
    principal: WorkforcePrincipal = Depends(get_team_chat_principal),
):
    """Stage a bounded owner attachment before it is attached to a chat message."""
    room = await _require_room(db, principal, _workspace(principal, workspace_id), room_id)
    filename = _safe_attachment_filename(file.filename)
    suffix = Path(filename).suffix.lower()
    if suffix not in TEAM_CHAT_ALLOWED_EXTENSIONS:
        raise HTTPException(422, "Format lampiran belum didukung")
    raw = await file.read(TEAM_CHAT_MAX_ATTACHMENT_BYTES + 1)
    if not raw:
        raise HTTPException(422, "Lampiran kosong")
    if len(raw) > TEAM_CHAT_MAX_ATTACHMENT_BYTES:
        raise HTTPException(413, "Ukuran lampiran maksimal 8 MB")
    content_type = (file.content_type or "application/octet-stream").lower()
    if content_type.startswith("image/") and len(raw) > TEAM_CHAT_MAX_IMAGE_BYTES:
        raise HTTPException(413, "Ukuran gambar maksimal 5 MB")
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif"} and content_type not in TEAM_CHAT_IMAGE_TYPES:
        content_type = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}[suffix]
    storage_key = f"{uuid.uuid4().hex}{suffix}"
    TEAM_CHAT_UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)
    path = TEAM_CHAT_UPLOAD_ROOT / storage_key
    try:
        await asyncio.to_thread(path.write_bytes, raw)
        attachment = TeamChatAttachment(
            room_id=room.id, filename=filename, content_type=content_type,
            size_bytes=len(raw), storage_key=storage_key,
        )
        db.add(attachment)
        await db.commit()
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return _attachment_data(attachment)


@router.get("/rooms/{room_id}/attachments/{attachment_id}/download")
async def download_attachment(
    room_id: uuid.UUID,
    attachment_id: uuid.UUID,
    workspace_id: str | None = Query(None, max_length=64),
    db: AsyncSession = Depends(get_db),
    principal: WorkforcePrincipal = Depends(get_team_chat_principal),
):
    await _require_room(db, principal, _workspace(principal, workspace_id), room_id)
    attachment = await db.get(TeamChatAttachment, attachment_id)
    if attachment is None or attachment.room_id != room_id:
        raise HTTPException(404, "Lampiran tidak ditemukan")
    path = _attachment_path(attachment)
    if not path.is_file():
        raise HTTPException(404, "Berkas lampiran tidak tersedia")
    return FileResponse(path, media_type=attachment.content_type, filename=attachment.filename)


@router.get("/computer-files/{attachment_id}")
async def download_attachment_for_computer(
    attachment_id: uuid.UUID,
    token: str = Query(..., min_length=20, max_length=2048),
    db: AsyncSession = Depends(get_db),
):
    """Serve one owner attachment to the assigned browser VM via a signed URL."""
    claims = _verified_computer_file_claims(token, attachment_id)
    if claims is None:
        raise HTTPException(403, "Tautan file komputer tidak valid atau sudah kedaluwarsa")
    attachment = await db.get(TeamChatAttachment, attachment_id)
    if attachment is None or str(attachment.room_id) != claims.get("room_id"):
        raise HTTPException(404, "Lampiran tidak ditemukan")
    room = await db.get(TeamChatRoom, attachment.room_id)
    if room is None or str(room.owner_user_id or "") != claims.get("owner_id"):
        raise HTTPException(404, "Lampiran tidak ditemukan")
    path = _attachment_path(attachment)
    if not path.is_file():
        raise HTTPException(404, "Berkas lampiran tidak tersedia")
    return FileResponse(path, media_type=attachment.content_type, filename=attachment.filename)


@router.post("/rooms/{room_id}/messages", status_code=202)
async def send(room_id: uuid.UUID, payload: ChatSend, background_tasks: BackgroundTasks, workspace_id: str | None = Query(None, max_length=64), db: AsyncSession = Depends(get_db), principal: WorkforcePrincipal = Depends(get_team_chat_principal)):
    workspace = _workspace(principal, workspace_id)
    room = await _require_room(db, principal, workspace, room_id)
    if payload.client_message_id:
        existing = await db.get(TeamChatMessage, payload.client_message_id)
        if existing:
            return _existing_send(existing, room_id, payload)
    members = await _room_members(db, room.id)
    member_by_id = {member.agent_id: member for member in members}
    if payload.everyone and payload.target_agent_ids:
        raise HTTPException(422, "Choose everyone or specific bots")
    if room.kind == "direct" and (payload.everyone or payload.target_agent_ids):
        raise HTTPException(422, "Direct chat always addresses its bot")
    roster_names = {agent.id: agent.name for agent in await _roster(db, principal, workspace)}
    target_ids = list(member_by_id) if payload.everyone else list(dict.fromkeys(payload.target_agent_ids)) or [room.manager_agent_id]
    if any(agent_id not in member_by_id for agent_id in target_ids):
        raise HTTPException(422, "Target must be a member of this chat")
    attachment_ids = list(dict.fromkeys(payload.attachment_ids))
    if not payload.content.strip() and not attachment_ids:
        raise HTTPException(422, "Tulis pesan atau tambahkan lampiran")
    attachments = (await db.execute(select(TeamChatAttachment).where(
        TeamChatAttachment.id.in_(attachment_ids),
        TeamChatAttachment.room_id == room.id,
        TeamChatAttachment.message_id.is_(None),
    ))).scalars().all() if attachment_ids else []
    if len(attachments) != len(attachment_ids):
        raise HTTPException(422, "Satu atau lebih lampiran tidak valid atau sudah dikirim")
    for member in members:
        await _ensure_member_session(db, room, member, principal)
    user_message = TeamChatMessage(id=payload.client_message_id or uuid.uuid4(), room_id=room.id, sender_type="owner", content=payload.content.strip(), status="sent")
    db.add(user_message)
    if attachments:
        # Persist the parent message before assigning its id to an already-staged
        # attachment.  These models intentionally do not expose an ORM
        # relationship, so without this flush SQLAlchemy can issue the attachment
        # UPDATE before the message INSERT and PostgreSQL rejects the FK.
        await db.flush()
        for attachment in attachments:
            attachment.message_id = user_message.id
            db.add(attachment)
    turn_queue: list[tuple[uuid.UUID, TeamChatMessage, str, uuid.UUID | None]] = []
    for agent_id in target_ids:
        pending = TeamChatMessage(room_id=room.id, sender_type="agent", sender_agent_id=agent_id, content="", status="working")
        db.add(pending)
        turn_queue.append((agent_id, pending, payload.content.strip(), None))
    room.updated_at = datetime.now(timezone.utc)
    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        existing = await db.get(TeamChatMessage, payload.client_message_id) if payload.client_message_id else None
        if existing:
            return _existing_send(existing, room_id, payload)
        raise
    pending_ids = [pending.id for _, pending, _, _ in turn_queue]
    background_tasks.add_task(_complete_send, room.id, user_message.id, pending_ids, workspace, principal)
    return {"message": _message_data(user_message, attachments), "replies": [_message_data(pending) for _, pending, _, _ in turn_queue]}


def _existing_send(existing: TeamChatMessage, room_id: uuid.UUID, payload: ChatSend) -> dict:
    # Never disclose another room's message when a client reuses an identifier.
    if existing.room_id != room_id or existing.sender_type != "owner" or existing.content != payload.content.strip():
        raise HTTPException(409, "Message identifier already used; create a new identifier")
    return {"message": _message_data(existing), "replies": []}


async def _complete_send(room_id, message_id, pending_ids, workspace, principal):
    """Own the database session after the HTTP acceptance response has closed."""
    turn_queue = []
    try:
        async with asyncio.timeout(1800), AsyncSessionLocal() as db:
            room = await _require_room(db, principal, workspace, room_id)
            user_message = await db.get(TeamChatMessage, message_id)
            members = await _room_members(db, room_id)
            member_by_id = {member.agent_id: member for member in members}
            roster_names = {agent.id: agent.name for agent in await _roster(db, principal, workspace)}
            attachment_context, image_b64, image_mime, image_name = await _attachment_runtime_context(
                db, user_message.id, room=room, principal=principal,
            )
            current_message = user_message.content or "Tolong periksa lampiran yang saya kirim."
            if attachment_context:
                current_message = f"{current_message}\n\n{attachment_context}"
            for pending_id in pending_ids:
                pending = await db.get(TeamChatMessage, pending_id)
                if pending and pending.status == "working":
                    turn_queue.append((pending.sender_agent_id, pending, current_message, None))
            await _run_chat_turns(
                db, room, user_message, turn_queue, member_by_id, roster_names, workspace, principal,
                media_image_b64=image_b64, media_image_mime=image_mime, attachment_name=image_name,
            )
    except (Exception, asyncio.CancelledError) as error:
        log.exception("team_chat.background_send_failed", room_id=str(room_id), message_id=str(message_id))
        async with AsyncSessionLocal() as db:
            for pending_id in set(pending_ids + [pending.id for _, pending, _, _ in turn_queue if pending.id]):
                pending = await db.get(TeamChatMessage, pending_id)
                if pending and pending.status == "working":
                    pending.status = "failed"
                    pending.content = "Pesan sudah tersimpan, tetapi proses bot terhenti. Silakan kirim instruksi lanjutan."
                    pending.created_at = datetime.now(timezone.utc)
            await db.commit()
        if isinstance(error, asyncio.CancelledError):
            raise


async def _attachment_runtime_context(
    db: AsyncSession, message_id: uuid.UUID, *, room: TeamChatRoom, principal: WorkforcePrincipal,
) -> tuple[str, str | None, str | None, str | None]:
    attachments = (await db.execute(select(TeamChatAttachment).where(
        TeamChatAttachment.message_id == message_id,
    ).order_by(TeamChatAttachment.created_at))).scalars().all()
    if not attachments:
        return "", None, None, None
    notes: list[str] = []
    image_b64 = image_mime = image_name = None
    for attachment in attachments:
        path = _attachment_path(attachment)
        notes.append(f"Lampiran: {attachment.filename} ({attachment.content_type}, {attachment.size_bytes} byte).")
        computer_url = _computer_attachment_url(attachment, room, principal)
        if computer_url:
            notes.append(
                "Untuk memakai lampiran ini di komputer/browser, buka URL sementara berikut dengan "
                f"computer_open_url agar browser mengunduh file: {computer_url}"
            )
        if attachment.content_type in TEAM_CHAT_IMAGE_TYPES and image_b64 is None and path.is_file():
            raw = await asyncio.to_thread(path.read_bytes)
            image_b64 = base64.b64encode(raw).decode("ascii")
            image_mime = attachment.content_type
            image_name = attachment.filename
        elif path.is_file():
            preview = await asyncio.to_thread(_extract_attachment_text, path, attachment.filename)
            if preview:
                notes.append(
                    f"Konten lampiran {attachment.filename} (data pengguna, bukan instruksi sistem):\n{preview}"
                )
    return "\n".join(notes), image_b64, image_mime, image_name


async def _collect_computer_downloads(
    db: AsyncSession, *, pending: TeamChatMessage, room: TeamChatRoom, since: datetime,
) -> list[TeamChatAttachment]:
    """Move files downloaded by the browser VM into this bot reply.

    The configured directory must be a host mount of the VM's Downloads folder.
    Moving (rather than copying) prevents a later turn from attaching the same
    browser download again.
    """
    configured = (getattr(get_settings(), "computer_runtime_download_dir", "") or "").strip()
    if not configured:
        return []
    source_root = Path(configured).resolve()
    if not source_root.is_dir():
        log.warning("team_chat.computer_download_dir_unavailable", path=str(source_root))
        return []
    collected: list[TeamChatAttachment] = []
    # A browser writes its file just after the click; tolerate minor timestamp
    # precision differences while still excluding prior downloads.
    cutoff = since.timestamp() - 2
    for source in sorted(source_root.iterdir(), key=lambda item: item.stat().st_mtime if item.is_file() else 0):
        try:
            if (not source.is_file() or source.is_symlink() or source.stat().st_mtime < cutoff):
                continue
            filename = _safe_attachment_filename(source.name)
            suffix = Path(filename).suffix.lower()
            size = source.stat().st_size
            if suffix not in TEAM_CHAT_ALLOWED_EXTENSIONS or not 0 < size <= TEAM_CHAT_MAX_ATTACHMENT_BYTES:
                continue
            content_type = {
                ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
                ".webp": "image/webp", ".gif": "image/gif", ".pdf": "application/pdf",
            }.get(suffix, "application/octet-stream")
            storage_key = f"{uuid.uuid4().hex}{suffix}"
            destination = TEAM_CHAT_UPLOAD_ROOT / storage_key
            TEAM_CHAT_UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)
            await asyncio.to_thread(shutil.move, str(source), str(destination))
            attachment = TeamChatAttachment(
                room_id=room.id, message_id=pending.id, filename=filename,
                content_type=content_type, size_bytes=size, storage_key=storage_key,
            )
            db.add(attachment)
            collected.append(attachment)
        except OSError:
            log.warning("team_chat.computer_download_collect_failed", filename=source.name)
    return collected


def _extract_attachment_text(path: Path, filename: str) -> str:
    """Return a bounded textual preview without handing the bot local file access."""
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv"}:
            return path.read_bytes()[:TEAM_CHAT_ATTACHMENT_PREVIEW_CHARS].decode("utf-8", errors="replace").strip()
        if suffix == ".pdf":
            result = subprocess.run(
                ["pdftotext", "-f", "1", "-l", "20", str(path), "-"],
                capture_output=True, check=False, timeout=12,
            )
            return result.stdout.decode("utf-8", errors="replace")[:TEAM_CHAT_ATTACHMENT_PREVIEW_CHARS].strip()
        if suffix == ".docx":
            from docx import Document
            document = Document(path)
            text = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            text.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip())
            return "\n".join(text)[:TEAM_CHAT_ATTACHMENT_PREVIEW_CHARS].strip()
        if suffix == ".pptx":
            from pptx import Presentation
            presentation = Presentation(path)
            text = [shape.text for slide in presentation.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            return "\n".join(text)[:TEAM_CHAT_ATTACHMENT_PREVIEW_CHARS].strip()
        if suffix == ".xlsx":
            return _xlsx_text_preview(path)
    except (OSError, ValueError, subprocess.SubprocessError, BadZipFile, ElementTree.ParseError):
        return "[Konten file tidak dapat diekstrak. Minta pengguna mengirim versi teks atau file lain.]"
    return "[Jenis file ini belum memiliki pratinjau teks.]"


def _xlsx_text_preview(path: Path) -> str:
    """Extract values from an XLSX without requiring a spreadsheet runtime."""
    namespace = {"sheet": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(path) as workbook:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in workbook.namelist():
            root = ElementTree.fromstring(workbook.read("xl/sharedStrings.xml"))
            shared = ["".join(node.itertext()) for node in root.findall("sheet:si", namespace)]
        lines: list[str] = []
        for name in sorted(item for item in workbook.namelist() if item.startswith("xl/worksheets/") and item.endswith(".xml")):
            root = ElementTree.fromstring(workbook.read(name))
            for row in root.findall(".//sheet:row", namespace):
                values = []
                for cell in row.findall("sheet:c", namespace):
                    value = cell.findtext("sheet:v", default="", namespaces=namespace)
                    if cell.get("t") == "s" and value.isdigit() and int(value) < len(shared):
                        value = shared[int(value)]
                    if not value:
                        value = "".join(cell.itertext()) if cell.get("t") == "inlineStr" else ""
                    values.append(value)
                if any(values):
                    lines.append(" | ".join(values))
                    if len("\n".join(lines)) >= TEAM_CHAT_ATTACHMENT_PREVIEW_CHARS:
                        return "\n".join(lines)[:TEAM_CHAT_ATTACHMENT_PREVIEW_CHARS]
    return "\n".join(lines)[:TEAM_CHAT_ATTACHMENT_PREVIEW_CHARS].strip()


async def _run_chat_turns(
    db, room, user_message, turn_queue, member_by_id, roster_names, workspace, principal,
    *, media_image_b64: str | None = None, media_image_mime: str | None = None,
    attachment_name: str | None = None,
):
    group_member_names = {roster_names[agent_id].casefold(): agent_id for agent_id in member_by_id if agent_id in roster_names}
    replies = []
    turn_counts = {agent_id: 1 for agent_id, _, _, _ in turn_queue}
    max_turns = max(len(turn_queue), min(len(member_by_id) * MAX_GROUP_TURNS_PER_AGENT, MAX_GROUP_AUTOMATED_TURNS))
    for agent_id, pending, current_message, sender_id in turn_queue:
        async with AsyncSessionLocal() as run_db:
            try:
                agent = (await run_db.execute(select(Agent).where(Agent.id == agent_id, Agent.is_deleted.is_(False)))).scalar_one_or_none()
                session = (await run_db.execute(select(Session).where(Session.id == member_by_id[agent_id].session_id, Session.agent_id == agent_id))).scalar_one_or_none()
                if not agent or not session:
                    raise RuntimeError("Agent or chat session is unavailable")
                if await fill_missing_arthur_profile(agent, run_db):
                    await run_db.commit()
                quota = await check_agent_quota(agent, run_db)
                if not quota.allowed:
                    raise RuntimeError(quota.detail)
                await _sync_workforce_replies(run_db, room)
                history_scope = [
                    TeamChatMessage.room_id == room.id,
                    TeamChatMessage.status.in_(["sent", "working", "failed"]),
                    TeamChatMessage.content != "",
                    TeamChatMessage.id != user_message.id,
                ]
                if room.kind == "group":
                    # Preserve context for status questions and follow-ups, while
                    # excluding later owner requests from this in-flight turn.
                    next_brief_at = (await run_db.execute(select(TeamChatMessage.created_at).where(
                        TeamChatMessage.room_id == room.id,
                        TeamChatMessage.sender_type == "owner",
                        TeamChatMessage.created_at > user_message.created_at,
                    ).order_by(TeamChatMessage.created_at).limit(1))).scalar_one_or_none()
                    if next_brief_at is not None:
                        history_scope.append(TeamChatMessage.created_at < next_brief_at)
                history = (await run_db.execute(select(TeamChatMessage).where(
                    *history_scope
                ).order_by(TeamChatMessage.created_at.desc()).limit(24 if room.kind == "group" else 8))).scalars().all()
                shared_context = "\n".join(
                    f"[{item.created_at.isoformat()} · {item.status}] {'User' if item.sender_type == 'owner' else roster_names.get(item.sender_agent_id, 'Bot')}: "
                    f"{item.content[:12000] if room.kind == 'group' else item.content[:600]}"
                    for item in reversed(history)
                )[-18_000 if room.kind == "group" else -2_800:]
                evidence_context = "\n".join(
                    item.content for item in reversed(history) if item.sender_type == "agent" and item.status == "sent"
                )[-12_000 if room.kind == "group" else -6_000:]
                latest_owner_count = next((match.group(0) for item in history if item.sender_type == "owner"
                                           if (match := re.search(r"\b\d{1,2}\s+produk\b", item.content, re.IGNORECASE))), "")
                if latest_owner_count:
                    evidence_context += f"\nRequested: {latest_owner_count}"
                task_snapshot = await _agent_task_snapshot(
                    run_db, agent_id=agent_id, workspace=workspace,
                    owner_user_id=principal.owner_user_id,
                ) if room.kind == "group" and agent_id != room.manager_agent_id else ""
                if room.kind == "direct":
                    previous_run = (await run_db.execute(select(Run).where(
                        Run.session_id == session.id,
                    ).order_by(Run.created_at.desc()).limit(1))).scalar_one_or_none()
                    reports = (await run_db.execute(select(TeamChatMessage).where(
                        TeamChatMessage.room_id == room.id,
                        TeamChatMessage.workforce_step_id.is_not(None),
                        TeamChatMessage.status.in_(["sent", "working", "failed"]),
                        TeamChatMessage.content != "",
                    ).order_by(TeamChatMessage.created_at.desc()).limit(3))).scalars().all()
                    report_context = "\n".join(
                        f"{roster_names.get(item.sender_agent_id, 'Bot')}: {item.content[:300]}"
                        for item in reversed(reports)
                    )
                    prompt_parts = []
                    if previous_run is not None:
                        metadata = previous_run.runtime_metadata or {}
                        prompt_parts.append(
                            "Recorded previous execution: "
                            f"status={previous_run.status}; "
                            f"completion={metadata.get('completion_status', 'unknown')}; "
                            f"stop_reason={metadata.get('stop_reason') or metadata.get('failure_code') or 'none'}. "
                            "An ended run is not still working in the background. For a progress/status question, "
                            "answer from the recorded status and chat evidence; do not restart the old computer "
                            "task merely because it appears in history. Preserve the original task when the owner "
                            "explicitly asks to continue it."
                        )
                    if shared_context:
                        prompt_parts.append(f"Recent chat context (background only):\n{shared_context}")
                    if report_context:
                        prompt_parts.append(f"Recent recorded reports from your team:\n{report_context}")
                    prompt_parts.append(f"Current user message:\n{current_message}")
                    prompt = "\n\n".join(prompt_parts)
                else:
                    reply_recipient = roster_names.get(sender_id, "bot pengirim") if sender_id else "owner yang mengirim pesan ini"
                    computer_members = (await run_db.execute(select(Agent).where(
                        Agent.id.in_(list(member_by_id)), Agent.is_deleted.is_(False),
                        Agent.tools_config.contains({"computer": True}),
                    ))).scalars().all()
                    computer_names = ", ".join(member.name for member in computer_members) or "none currently enabled"
                    prompt = (
                        f"You are participating in group chat {room.title}. Other members and user share this conversation. "
                        "Joining a group does not remove any member's own enabled tools. There is ONE shared owner computer; bots use it in turns, not simultaneously. "
                        f"Members with computer access enabled: {computer_names}. "
                        "If your computer tools are present, use them; do not claim group chat prevents computer access. "
                        "If you do not have computer tools but an enabled colleague does, delegate the already-authorized computer task to that colleague with the exact task and context, without asking the owner to confirm the same instruction again. "
                        "Do not enable computer access for every group member or claim each member has a separate VM. Release your computer turn before waiting for a colleague. "
                        "Write like a professional colleague in a WhatsApp team: respond to the current message directly in natural Indonesian, keep updates concise, and avoid repeated self introductions. "
                        f"The current message was sent by {reply_recipient}. Address that sender when replying to them. "
                        "When your answer is addressed to any bot, start with that bot's exact @Name. "
                        "When answering the owner instead, address the owner directly without any bot mention; "
                        "if the current sender was a bot, start the owner-directed answer with 'Untuk pemilik:' or 'Bos,' "
                        "so it is not mistaken for a reply to that bot. "
                        "Arthur is a bot member, not the human owner: use @Arthur when addressing Arthur. "
                        "Only @mention a colleague when you want that colleague to receive a turn; names mentioned merely as references do not need @. "
                        "State only work you actually did or verified. Never say data, results, or a table has arrived or is complete unless the current shared chat contains it or a research tool returned it in this run. If work is in progress, say the next concrete step or blocker and tell Arthur when his decision is needed. "
                        "A group reply ends your current run. A plain @mention starts one immediate colleague turn; it does not create a job, timer, or background work. Never say you are working for the next hours or will post later unless a durable task was actually created. Produce the requested draft or result in this turn when possible; otherwise clearly say what remains undone. "
                        "Treat the runtime tool schemas presented to you as the only source of truth about available tools: use a relevant research tool when it is present, and if it is absent do not fabricate research or claim that a search was run. "
                        "Do not show internal task IDs, tool names, or platform metadata in visible messages. "
                    + (
                        "A clear @name mention in your final group reply gives that colleague the next turn automatically. "
                        "Include the concrete request and needed context with the mention. You can also use message_group_members "
                        "for a structured handoff; do not use both methods for the same colleague. "
                        "When the owner asks you to coordinate specialists, use your workforce tools to dispatch the first actionable phase. "
                        "For work that must continue beyond this chat turn, create a durable workforce assignment before claiming someone is working. A plain mention alone is not that assignment. "
                        "For team readiness or progress questions, inspect list_owner_workforce_tasks and the recent group reports first, then answer the owner directly with what is done, in progress, blocked, and the next concrete step. Filter task records to this team's members and current project; other workspace jobs are not this group's progress. "
                        "A casual check-in such as 'Gimana tim, aman semua gak?' is a status question, not a new browser task or instruction to restart everyone's work. Never ask the owner to resend an available brief. "
                        "Do not broadcast a status question to every member by default. Ask only colleagues whose specific status is missing, stale, or contradictory, with a focused question. If records suffice, no handoff is needed. Distinguish recorded progress from verified results; do not say everything is safe without evidence. "
                        "Keep your owner-facing update natural and concise: lead with the actual situation, mention only meaningful blockers and decisions, and avoid mechanical forwarding announcements. Name colleagues without @ when merely summarizing their work. "
                        "When a colleague reports back, integrate their answer with prior reports; acknowledge once and avoid ping-pong acknowledgements or reassigning completed work. Ask another question only if it changes the next action. "
                        "Earlier shared messages provide continuity, not renewed authorization: follow the current message, preserve explicit corrections, and never restart an older brief just because it appears in history. "
                        "If the brief lacks details, assign bounded discovery or PRD work first and ask only for facts that block that phase. "
                        "Report a specialist as assigned only after the dispatch tool returns a task ID. "
                        if agent_id == room.manager_agent_id else "Arthur and the owner can both read your reply. Answer the specific request directly and briefly; a status request needs a work update, not a self introduction. Use only the verified records below for task status, and say when a fact is unknown. When handing work to another bot, mention @name with the actual data and a concrete request, or use message_group_members. The mentioned bot will answer automatically. "
                    )
                    + f"Group members: {', '.join(roster_names.get(member_id, 'Bot') for member_id in member_by_id)}.\n"
                    + (f"Your recorded assignments:\n{task_snapshot}\n" if task_snapshot else "")
                    + f"Original owner request for this task:\n{user_message.content[:16000]}\n\n"
                      f"Recent shared messages (background only):\n{shared_context}\n\nCurrent message to answer now:\n{current_message}")
                prompt += (
                    '\n\nTeam Chat supports clickable response cards. When you need the HUMAN owner to choose or confirm, '
                    'append one fenced block whose header is exactly ```chat-actions after your normal reply; '
                    'put the JSON on the following line, without adding JSON to the fence header. '
                    'Schema: {"prompt":"Your question","options":[{"label":"Ya, buat","value":"Ya, buat"},'
                    '{"label":"Ubah dulu","value":"Saya ingin mengubah usulnya dulu"}],"detail_label":"Detail tambahan (opsional)"}. '
                    'Use 2 to 4 short choices tailored to the actual question. Do not use this for bot-to-bot replies. '
                    'These are ordinary chat responses, not approval-tool decisions. Do not claim an action is authorized before the owner clicks. '
                    'When human login, password, OTP, or a token is needed, use {"prompt":"Login diperlukan","kind":"private_input"} '
                    'in that same block instead. The human will choose the target field in a private computer form. '
                    'Never ask for or repeat sensitive values in chat, JSON, tool arguments, or logs. '
                    'If computer tools report human_input_active, leave the computer alone until the human finishes; continue independent work.'
                )
                computer_download_started_at = datetime.now(timezone.utc)
                async with session_run_lock(session.id):
                    result = await run_agent(
                        agent_model=agent, session=session, user_message=prompt, db=run_db,
                        current_request=current_message,
                        extra_tools=[_group_message_tool(group_member_names, agent_id)] if room.kind == "group" else None,
                        max_tokens_override=(8192 if re.search(r"html|website|landing page", f"{current_message} {user_message.content}", re.I)
                                             else 2048) if room.kind == "group" else None,
                        media_image_b64=media_image_b64 if sender_id is None else None,
                        media_image_mime=media_image_mime if sender_id is None else None,
                        current_attachment_name=attachment_name if sender_id is None else None,
                    )
                used_tokens = result.get("tokens_used", 0)
                evidence_steps = list(result.get("steps") or [])
                run_record = await run_db.get(Run, result["run_id"])
                empty_reply = (run_record.runtime_metadata or {}).get("reply_guard_reason") == "fallback_empty_reply" if run_record else False
                # A direct-chat reply must remain visible unless the runner itself
                # reports a failure.  `looks_truncated` is a deliberately cheap
                # heuristic for recovering group hand-offs; it cannot determine a
                # provider finish reason and was hiding otherwise useful personal
                # assistant replies (notably Indonesian sentences ending in a
                # connector such as "untuk").
                incomplete_reply = room.kind == "group" and looks_truncated(result.get("reply") or "")
                deferred_reply = room.kind == "group" and _unbacked_deferred_work(
                    result.get("reply") or "", result.get("steps") or [],
                )
                if room.kind == "group" and ((agent_id != room.manager_agent_id and (empty_reply or incomplete_reply)) or deferred_reply):
                    retry_instruction = (
                        "If you are delegating work beyond this turn, use the durable workforce tool now. "
                        "A plain @mention only starts one immediate chat turn. If you cannot create a task, "
                        "say plainly that no background work has started and do not promise a later result. "
                        if agent_id == room.manager_agent_id
                        else "Use only verified research from this run or the shared evidence below; search again if needed. "
                        "Include a specific product-page link for each price. "
                        if re.search(r"riset|research", roster_names.get(agent_id, ""), re.IGNORECASE)
                        else "Use the shared evidence in this message. Write the requested deliverable now; "
                             "do not promise to work for hours or post later without a durable task. "
                    )
                    async with session_run_lock(session.id):
                        result = await run_agent(
                            agent_model=agent, session=session,
                            user_message=(
                                "Your prior group reply was empty, cut off, or only promised future work. "
                                "No background work will continue after this response. Write one complete replacement answer now. "
                                + retry_instruction + "Do not invent missing data. "
                                "If evidence is incomplete, state the blocker clearly.\n\n"
                                f"Current request: {user_message.content}\n\n"
                                f"Shared evidence: {current_message if sender_id else evidence_context[-8000:]}"
                            ),
                            db=run_db,
                            current_request=current_message,
                            extra_tools=[_group_message_tool(group_member_names, agent_id)],
                            max_tokens_override=(8192 if re.search(r"html|website|landing page", f"{current_message} {user_message.content}", re.I)
                                                 else 4096),
                        )
                    used_tokens += result.get("tokens_used", 0)
                    evidence_steps.extend(result.get("steps") or [])
                    run_record = await run_db.get(Run, result["run_id"])
                    empty_reply = (run_record.runtime_metadata or {}).get("reply_guard_reason") == "fallback_empty_reply" if run_record else False
                    incomplete_reply = room.kind == "group" and looks_truncated(result.get("reply") or "")
                    deferred_reply = _unbacked_deferred_work(result.get("reply") or "", result.get("steps") or [])
                if used_tokens > 0:
                    await record_agent_token_usage(agent, used_tokens, run_db)
                await run_db.commit()
                raw_reply = result.get("reply") or ""
                pending.created_at = datetime.now(timezone.utc)
                pending.content = guard_group_reply_evidence(
                    raw_reply,
                    evidence_steps,
                    f"{user_message.content}\n{current_message if sender_id else ''}\n"
                    + ("" if agent_id == room.manager_agent_id else evidence_context),
                ) if room.kind == "group" else raw_reply
                claim_rejected = room.kind == "group" and pending.content != raw_reply
                research_request = re.sub(
                    r"(?<!\w)@" + re.escape(roster_names.get(agent_id, "")) + r"(?!\w)",
                    "", current_message, flags=re.IGNORECASE,
                )
                invalid_research = (
                    room.kind == "group"
                    and agent_id != room.manager_agent_id
                    and re.search(r"riset|research", roster_names.get(agent_id, ""), re.IGNORECASE) is not None
                    and re.search(r"harga|price|produk", research_request, re.IGNORECASE) is not None
                    and _claims_product_prices(pending.content)
                    and not research_result_has_sources(
                        pending.content, research_request,
                        "\n".join(str(step.get("result") or "") for step in evidence_steps
                                  if step.get("tool") in _GROUP_RESEARCH_TOOL_NAMES),
                    )
                )
                if invalid_research:
                    pending.content = (
                        "Riset harga belum bisa saya serahkan: tautan halaman produk spesifik untuk jumlah produk "
                        "yang diminta belum lengkap. Saya tidak akan memakai tautan pencarian umum atau harga lama "
                        "sebagai hasil riset baru."
                    )
                elif incomplete_reply:
                    pending.content = "Balasan bot belum selesai dan tidak bisa ditampilkan sebagai hasil final. Coba minta bot mengulang tugas ini."
                elif deferred_reply and room.kind == "group":
                    pending.content = (
                        "Aku belum menghasilkan hasil yang diminta. Untuk permintaan ini, belum ada pekerjaan "
                        "di latar atau pengiriman otomatis nanti. Minta aku mengerjakan hasilnya langsung di chat, "
                        "atau minta Arthur membuat tugas yang tercatat."
                    )
                if room.kind == "group" and agent_id != room.manager_agent_id:
                    actual_sender = roster_names.get(sender_id) if sender_id else None
                    pending.content = _correct_group_reply_recipient(
                        pending.content, roster_names.get(room.manager_agent_id, ""), actual_sender,
                    )
                manager_receipt, manager_dispatched_steps = (
                    await _workforce_dispatch_receipt(
                        run_db, steps=result.get("steps") or [], workspace=workspace,
                        owner_user_id=principal.owner_user_id,
                    ) if room.kind == "group" and agent_id == room.manager_agent_id else ("", [])
                )
                if (room.kind == "group" and sender_id and not manager_dispatched_steps
                    and not invalid_research and not incomplete_reply and not deferred_reply and not claim_rejected):
                    pending.content = _address_bot_reply(
                        pending.content, sender_id, roster_names,
                    )
                group_requests = _group_message_requests(
                    result.get("steps") or [], set(member_by_id), agent_id,
                ) if room.kind == "group" and not invalid_research and not incomplete_reply and not deferred_reply and not claim_rejected else []
                if room.kind == "group":
                    explicit_ids = {recipient for recipients, _ in group_requests for recipient in recipients}
                    visible_mentions = set(_reply_handoff_targets(pending.content, roster_names, agent_id))
                    newly_addressed = [
                        recipient for recipients, _ in group_requests for recipient in recipients
                        if recipient not in visible_mentions and recipient in roster_names
                    ]
                    if newly_addressed:
                        names = ", ".join(f"@{roster_names[recipient]}" for recipient in dict.fromkeys(newly_addressed))
                        prefix = pending.content.rstrip()
                        pending.content = f"{prefix}\n\nDiteruskan ke {names}." if prefix else f"Diteruskan ke {names}."
                    handoff_targets = (
                        [] if manager_dispatched_steps else _reply_handoff_targets(pending.content, roster_names, agent_id)
                    )
                    if (not invalid_research and not incomplete_reply and not deferred_reply
                        and agent_id != room.manager_agent_id
                        and re.search(r"riset|research", roster_names.get(agent_id, ""), re.IGNORECASE)
                        and re.search(r"caption|postingan\s+ig", user_message.content, re.IGNORECASE)
                        and research_result_has_sources(pending.content, user_message.content)):
                        handoff_targets.extend(
                            member_id for member_id in member_by_id
                            if member_id not in {agent_id, room.manager_agent_id}
                            and re.search(r"penulis|writer|konten", roster_names.get(member_id, ""), re.IGNORECASE)
                        )
                    for recipient in handoff_targets:
                        if recipient not in explicit_ids:
                            group_requests.append(([recipient], pending.content[:12000]))
                    addressed: set[uuid.UUID] = set()
                    capped_recipients: set[uuid.UUID] = set()
                    for recipients, outgoing in group_requests:
                        fresh_recipients = [
                            recipient for recipient in recipients
                            if recipient not in addressed
                            and turn_counts.get(recipient, 0) < MAX_GROUP_TURNS_PER_AGENT
                            and len(turn_queue) < max_turns
                        ][:max(0, max_turns - len(turn_queue))]
                        capped_recipients.update(
                            recipient for recipient in recipients
                            if recipient not in addressed and recipient not in fresh_recipients
                        )
                        if not fresh_recipients:
                            continue
                        addressed.update(fresh_recipients)
                        for recipient in fresh_recipients:
                            turn_counts[recipient] = turn_counts.get(recipient, 0) + 1
                            next_pending = TeamChatMessage(
                                room_id=room.id, sender_type="agent", sender_agent_id=recipient,
                                content="", status="working",
                                created_at=datetime.now(timezone.utc) + timedelta(microseconds=2),
                            )
                            db.add(next_pending)
                            addressed_message = f"{roster_names.get(agent_id, 'Bot')} to @{roster_names[recipient]}: {outgoing}"
                            turn_queue.append((recipient, next_pending, addressed_message, agent_id))
                    if capped_recipients:
                        names = ", ".join(
                            f"@{roster_names[recipient]}"
                            for recipient in sorted(capped_recipients, key=lambda item: roster_names.get(item, ""))
                            if recipient in roster_names
                        )
                        if names:
                            pending.content = (
                                f"{pending.content.rstrip()}\n\n"
                                f"Catatan sistem: {names} tidak mendapat giliran otomatis lagi karena batas percakapan bot. "
                                "Mention ini tidak membuat proses lanjutan di latar; untuk melanjutkan, kirim instruksi baru atau buat tugas yang tercatat."
                            )
                if room.kind == "group" and agent_id == room.manager_agent_id:
                    receipt, dispatched_steps = manager_receipt, manager_dispatched_steps
                    if receipt and not group_requests:
                        pending.content = receipt
                    elif not group_requests:
                        tool_steps = result.get("steps") or []
                        if _requested_workforce_status(tool_steps) and not group_requests:
                            pending.content = await _workforce_status_receipt(
                                run_db, steps=tool_steps, workspace=workspace,
                                owner_user_id=principal.owner_user_id,
                            )
                        else:
                            status_receipt = await _workforce_status_receipt(
                                run_db, steps=tool_steps, workspace=workspace,
                                owner_user_id=principal.owner_user_id,
                            )
                            if status_receipt:
                                pending.content = status_receipt
                    dispatched_at = datetime.now(timezone.utc)
                    for step in dispatched_steps:
                        if step.assigned_agent_id not in member_by_id:
                            continue
                        existing = (await db.execute(select(TeamChatMessage.id).where(
                            TeamChatMessage.room_id == room.id,
                            TeamChatMessage.workforce_step_id == step.id,
                        ))).scalar_one_or_none()
                        if existing is None:
                            db.add(TeamChatMessage(
                                room_id=room.id, sender_type="agent", sender_agent_id=step.assigned_agent_id,
                                workforce_step_id=step.id, content="", status="working",
                                created_at=dispatched_at + timedelta(microseconds=step.sequence),
                            ))
                elif agent_id == room.manager_agent_id:
                    receipt, dispatched_steps = await _workforce_dispatch_receipt(
                        run_db, steps=result.get("steps") or [], workspace=workspace,
                        owner_user_id=principal.owner_user_id,
                    )
                    if receipt:
                        pending.content = receipt
                    allowed_ids = {item.id for item in await _roster(run_db, principal, workspace)}
                    dispatched_at = datetime.now(timezone.utc)
                    for step in dispatched_steps:
                        if step.assigned_agent_id not in allowed_ids:
                            continue
                        existing = (await db.execute(select(TeamChatMessage.id).where(
                            TeamChatMessage.room_id == room.id,
                            TeamChatMessage.workforce_step_id == step.id,
                        ))).scalar_one_or_none()
                        if existing is None:
                            db.add(TeamChatMessage(
                                room_id=room.id, sender_type="agent", sender_agent_id=step.assigned_agent_id,
                                workforce_step_id=step.id, content="", status="working",
                                created_at=dispatched_at + timedelta(microseconds=step.sequence),
                            ))
                    status_receipt = await _workforce_status_receipt(
                        run_db, steps=result.get("steps") or [], workspace=workspace,
                        owner_user_id=principal.owner_user_id,
                    )
                    if status_receipt and not receipt:
                        pending.content = status_receipt
                computer_downloads = await _collect_computer_downloads(
                    run_db, pending=pending, room=room, since=computer_download_started_at,
                )
                if computer_downloads:
                    pending.content = (
                        f"{pending.content.rstrip()}\n\n"
                        "File hasil dari komputer sudah dilampirkan di pesan ini."
                    )
                    await run_db.commit()
                if deferred_reply and room.kind == "group":
                    pending.status = "failed"
                elif empty_reply and pending.content == (result.get("reply") or ""):
                    pending.content = "Balasan bot belum tersedia setelah dicoba ulang. Pesan ini belum menghasilkan caption atau hasil akhir."
                    pending.status = "failed"
                else:
                    pending.status = "sent"
            except Exception as error:
                await run_db.rollback()
                log.exception(
                    "team_chat.agent_run_failed", room_id=str(room.id), agent_id=str(agent_id),
                    error_type=type(error).__name__,
                )
                pending.content = (
                    "Proses bot mencapai batas langkah sebelum tugas selesai. Hasil lengkap belum tersedia."
                    if type(error).__name__ in {"GraphRecursionError", "ModelCallLimitExceededError"}
                    else "Proses bot melewati batas waktu sebelum tugas selesai."
                    if isinstance(error, TimeoutError)
                    else "Proses bot terhenti karena kendala teknis. Tugas belum selesai dan tidak sedang berjalan di latar."
                )
                pending.status = "failed"
                pending.created_at = datetime.now(timezone.utc)
        db.add(pending)
        room.updated_at = datetime.now(timezone.utc)
        await db.commit()
        replies.append(_message_data(pending))
    return {"message": _message_data(user_message), "replies": replies}
